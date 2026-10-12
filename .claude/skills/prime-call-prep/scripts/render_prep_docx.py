"""Deterministic Call Prep renderer: markdown -> docx using template styles.

Usage: python3 render_prep_docx.py input.md "Output Name.docx"

BRANDING: edit the five constants below to match your brand.
Defaults are the Prime AI palette: black, white, gold.
Keep heading colors dark enough to read on white paper.
"""
import os, re, sys, copy
from docx import Document
from docx.shared import Pt, Inches
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

# template.docx ships with this skill in assets/; falls back to a blank document
TEMPLATE = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "assets", "template.docx")

# Brand constants (EDIT THESE for your own brand)
GOLD = "A67C00"        # heading accent, WCAG-safe on white at heading sizes
GOLD_BORDER = "C9A227" # borders and accents
BLACK = "1A1A1A"       # body text
HEAD_FONT = "Georgia"
BODY_FONT = "Calibri"

def apply_brand(doc):
    from docx.shared import RGBColor
    for s in doc.styles:
        try:
            name = s.name
        except Exception:
            continue
        if name == 'Title':
            s.font.name = HEAD_FONT; s.font.size = Pt(21); s.font.bold = True
            s.font.color.rgb = RGBColor.from_string(BLACK)
        elif name in ('Heading 1', 'Heading 2', 'Heading 3'):
            s.font.name = HEAD_FONT; s.font.bold = True
            s.font.color.rgb = RGBColor.from_string(GOLD)
            if name == 'Heading 1': s.font.size = Pt(15)
        elif name in ('Normal', 'List Paragraph'):
            s.font.name = BODY_FONT
            s.font.color.rgb = RGBColor.from_string(BLACK)

def shade_cell(cell, hex_fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear'); shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), hex_fill)
    tcPr.append(shd)

def get_style(doc, name):
    for s in doc.styles:
        if s.name == name:
            return s
    return None

def make_bullet(paragraph):
    """Attach the template's bullet numbering (numId 1, level 0)."""
    pPr = paragraph._p.get_or_add_pPr()
    numPr = OxmlElement('w:numPr')
    ilvl = OxmlElement('w:ilvl'); ilvl.set(qn('w:val'), '0')
    numId = OxmlElement('w:numId'); numId.set(qn('w:val'), '1')
    numPr.append(ilvl); numPr.append(numId)
    pPr.append(numPr)
    paragraph.paragraph_format.left_indent = Inches(0.25)
    paragraph.paragraph_format.space_after = Pt(6)

def add_hyperlink(paragraph, text, url):
    """Insert a real clickable hyperlink run."""
    from docx.opc.constants import RELATIONSHIP_TYPE as RT
    part = paragraph.part
    r_id = part.relate_to(url, RT.HYPERLINK, is_external=True)
    hyperlink = OxmlElement('w:hyperlink')
    hyperlink.set(qn('r:id'), r_id)
    run = OxmlElement('w:r')
    rPr = OxmlElement('w:rPr')
    rf = OxmlElement('w:rFonts')
    for attr in ('w:ascii', 'w:hAnsi', 'w:cs'):
        rf.set(qn(attr), BODY_FONT)
    color = OxmlElement('w:color'); color.set(qn('w:val'), GOLD)
    u = OxmlElement('w:u'); u.set(qn('w:val'), 'single')
    rPr.append(rf); rPr.append(color); rPr.append(u)
    run.append(rPr)
    t = OxmlElement('w:t'); t.text = text
    t.set(qn('xml:space'), 'preserve')
    run.append(t)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)

def _body_font(run):
    run.font.name = BODY_FONT
    rPr = run._element.get_or_add_rPr()
    rf = rPr.find(qn('w:rFonts'))
    if rf is None:
        rf = OxmlElement('w:rFonts'); rPr.append(rf)
    for attr in ('w:ascii', 'w:hAnsi', 'w:cs'):
        rf.set(qn(attr), BODY_FONT)

def add_fmt(paragraph, text):
    """Add runs handling **bold** and *italic* (no links). Forces body font on every run."""
    pos = 0
    for m in re.finditer(r'\*\*(.+?)\*\*|\*(.+?)\*', text):
        if m.start() > pos:
            _body_font(paragraph.add_run(text[pos:m.start()]))
        if m.group(1) is not None:
            r = paragraph.add_run(m.group(1)); r.bold = True; _body_font(r)
        else:
            r = paragraph.add_run(m.group(2)); r.italic = True; _body_font(r)
        pos = m.end()
    if pos < len(text):
        _body_font(paragraph.add_run(text[pos:]))

def add_runs(paragraph, text):
    """Add runs, handling [text](url) links first, then **bold** and *italic*."""
    pos = 0
    for m in re.finditer(r'\[([^\]]+)\]\((https?://[^)\s]+)\)', text):
        if m.start() > pos:
            add_fmt(paragraph, text[pos:m.start()])
        add_hyperlink(paragraph, m.group(1), m.group(2))
        pos = m.end()
    if pos < len(text):
        add_fmt(paragraph, text[pos:])

def set_cell_borders(cell, color=GOLD_BORDER, sz=8):
    tcPr = cell._tc.get_or_add_tcPr()
    borders = OxmlElement('w:tcBorders')
    for edge in ('top', 'left', 'bottom', 'right'):
        el = OxmlElement(f'w:{edge}')
        el.set(qn('w:val'), 'single'); el.set(qn('w:sz'), str(sz))
        el.set(qn('w:space'), '0'); el.set(qn('w:color'), color)
        borders.append(el)
    tcPr.append(borders)

def parse_md(path):
    lines = open(path).read().split('\n')
    blocks, i = [], 0
    while i < len(lines):
        ln = lines[i]
        if ln.startswith('# ') and not ln.startswith('## '):
            blocks.append(('title', ln[2:].strip())); i += 1
        elif ln.startswith('## '):
            blocks.append(('heading', ln[3:].strip())); i += 1
        elif ln.startswith('+--'):
            # grid table box: collect | lines until closing +--
            i += 1; cell_lines = []
            while i < len(lines) and not lines[i].startswith('+--'):
                raw = lines[i]
                if raw.startswith('|'):
                    inner = raw.strip().strip('|').strip()
                    cell_lines.append(inner)
                i += 1
            i += 1  # skip closing fence
            paras, cur = [], []
            for cl in cell_lines:
                if cl == '':
                    if cur: paras.append(' '.join(cur)); cur = []
                else:
                    cur.append(cl)
            if cur: paras.append(' '.join(cur))
            blocks.append(('box', paras))
        elif ln.startswith('- '):
            item = [ln[2:].strip()]
            i += 1
            while i < len(lines) and lines[i].startswith(('  ', '\t')) and lines[i].strip():
                item.append(lines[i].strip()); i += 1
            blocks.append(('bullet', ' '.join(item)))
        elif ln.startswith('|'):
            rows = []
            while i < len(lines) and lines[i].startswith('|'):
                cells = [c.strip() for c in lines[i].strip().strip('|').split('|')]
                if not all(re.fullmatch(r':?-{2,}:?', c) for c in cells):
                    rows.append(cells)
                i += 1
            blocks.append(('table', rows))
        elif ln.strip():
            blocks.append(('para', ln.strip())); i += 1
        else:
            i += 1
    return blocks

def render(md_path, out_path):
    doc = Document(TEMPLATE) if os.path.exists(TEMPLATE) else Document()
    apply_brand(doc)
    # wipe template body
    body = doc.element.body
    for el in list(body):
        if el.tag != qn('w:sectPr'):
            body.remove(el)
    blocks = parse_md(md_path)
    n = len(blocks)
    for idx, (kind, data) in enumerate(blocks):
        if kind == 'title':
            p = doc.add_paragraph(style=get_style(doc, 'Title')); add_runs(p, data)
        elif kind == 'heading':
            p = doc.add_paragraph(style=get_style(doc, 'Heading 1')); add_runs(p, data)
        elif kind == 'box':
            t = doc.add_table(rows=1, cols=1)
            cell = t.cell(0, 0)
            set_cell_borders(cell)
            cell.paragraphs[0].text = ''
            first = True
            for ptext in data:
                p = cell.paragraphs[0] if first else cell.add_paragraph()
                first = False
                add_runs(p, ptext)
                p.paragraph_format.space_after = Pt(6)
                p.paragraph_format.space_before = Pt(2)
            doc.add_paragraph()  # spacer
        elif kind == 'bullet':
            p = doc.add_paragraph(style=get_style(doc, 'List Paragraph'))
            make_bullet(p)
            add_runs(p, data)
        elif kind == 'table':
            header, rows = data[0], data[1:]
            t = doc.add_table(rows=len(rows) + 1, cols=len(header))
            t.autofit = False
            t.allow_autofit = False
            if len(header) == 2:
                widths = [Inches(1.7), Inches(4.8)]
            elif len(header) == 3:
                widths = [Inches(1.5), Inches(0.7), Inches(4.3)]
            else:
                widths = None
            if widths:
                tbl = t._tbl
                tblGrid = tbl.find(qn('w:tblGrid'))
                if tblGrid is not None:
                    for gc, w in zip(tblGrid.findall(qn('w:gridCol')), widths):
                        gc.set(qn('w:w'), str(int(w.emu / 635)))
            from docx.shared import RGBColor
            for c, htext in enumerate(header):
                cell = t.cell(0, c); set_cell_borders(cell, sz=6)
                shade_cell(cell, BLACK)
                p = cell.paragraphs[0]
                r = p.add_run(re.sub(r'\*\*(.+?)\*\*', r'\1', htext)); r.bold = True
                r.font.color.rgb = RGBColor.from_string('FFFFFF')
                _body_font(r)
                if widths: cell.width = widths[c]
            for ri, row in enumerate(rows, start=1):
                for c, ctext in enumerate(row[:len(header)]):
                    cell = t.cell(ri, c); set_cell_borders(cell, sz=6)
                    add_runs(cell.paragraphs[0], ctext)
                    if widths: cell.width = widths[c]
            doc.add_paragraph()
        elif kind == 'para':
            p = doc.add_paragraph(); add_runs(p, data)
            if idx == n - 1:  # footer line already italic via * *
                pass
    doc.save(out_path)
    print("rendered", out_path)

if __name__ == '__main__':
    render(sys.argv[1], sys.argv[2])
